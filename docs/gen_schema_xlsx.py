# -*- coding: utf-8 -*-
"""현재 SQLite(app.db) 스키마를 Excel(.xlsx)로 정리 출력하는 유틸 스크립트.

실행: python docs/gen_schema_xlsx.py  (프로젝트 루트 기준)
필요: pip install openpyxl
스키마 변경 시 재실행하면 docs/DB_스키마.xlsx 가 갱신된다.
"""
import sqlite3
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 이 파일(docs/gen_schema_xlsx.py)의 상위 = 프로젝트 루트
ROOT = Path(__file__).resolve().parent.parent
DB = ROOT / "data" / "app.db"
OUT = ROOT / "docs" / "DB_스키마.xlsx"

TABLE_DESC = {
    "all_stock_master": "전 증시 종목 마스터 (KOSPI/KOSDAQ/ETF 종목 기본정보)",
    "investor_trading_daily": "종목별·일자별 OHLCV 가격 + 주체별 순매수 수급",
    "market_indices_daily": "코스피·코스닥·S&P500 지수 및 원/달러 환율 일별 데이터",
    "paper_portfolios": "모의투자(rec)/투자제안(prop) 자산 계좌 상태 (계좌유형 분리)",
    "paper_positions": "모의투자/투자제안 보유 종목 잔고",
    "paper_trade_histories": "모의투자/투자제안 매수·매도 체결 로그",
    "strategy_daily_equity": "전략(combo)별 일별 누적 평가자산 (Equity Curve)",
    "strategy_leaderboard": "전략 조합별 백테스트 성과 리더보드 (수익률·승률·MDD 등)",
    "strategy_trade_logs": "백테스트 매매일지 / 개별 매수·매도 체결 로그",
    "sync_logs": "데이터 수집·동기화 파이프라인 실행 이력",
    "target_stocks": "수집 대상 타깃 종목 목록 (KOSPI200 + KOSDAQ150 + 미국ETF)",
}

COL_DESC = {
    "all_stock_master": {
        "code": "종목코드 (PK)", "name": "종목명", "market": "시장구분 (KOSPI/KOSDAQ/ETF 등)",
        "industry": "세부 산업분류 (예: 반도체와반도체장비)",
        "sector": "지수구분 (KOSPI 200 / KOSDAQ 150 / ETF_USA / 일반)",
        "marcap": "시가총액 (원)", "stocks": "상장주식수",
        "created_at": "등록일시", "updated_at": "수정일시",
    },
    "investor_trading_daily": {
        "symbol": "종목코드 (복합 PK)", "date": "일자 YYYYMMDD/YYYY-MM-DD (복합 PK)",
        "personal_net_buy": "개인 순매수 (원)", "foreigner_net_buy": "외국인 순매수 (원)",
        "institution_net_buy": "기관 순매수 (원)", "pension_net_buy": "연기금 순매수 (원)",
        "financial_net_buy": "금융투자 순매수 (원)", "other_corp_net_buy": "기타법인 순매수 (원)",
        "close_price": "종가 (원)", "open_price": "시가 (원)", "high_price": "고가 (원)",
        "low_price": "저가 (원)", "volume": "거래량 (주)", "updated_at": "수정일시",
        "is_suspended": "거래정지 여부 (1=정지). 판별식: volume=0 AND high=low",
    },
    "market_indices_daily": {
        "date": "일자 (PK)", "kospi_close": "코스피 지수 종가", "kosdaq_close": "코스닥 지수 종가",
        "sp500_close": "S&P500 지수 종가", "usdkrw_rate": "원/달러 환율", "updated_at": "수정일시",
    },
    "paper_portfolios": {
        "id": "자산 계좌 식별 ID (PK)", "initial_balance": "초기 투자 자산 (원)",
        "cash_balance": "현재 잔여 현금 (원)", "total_asset_value": "총 자산 평가액 (원)",
        "updated_at": "최종 업데이트 일시",
        "account_type": "계좌 유형 (rec: 모의투자, prop: 투자제안)",
    },
    "paper_positions": {
        "id": "포지션 ID (PK)", "stock_code": "종목 코드", "stock_name": "종목명",
        "buy_date": "매수일 (YYYY-MM-DD)", "buy_price": "매수가 (당일 종가)",
        "quantity": "보유 수량", "total_amount": "총 투입 금액 (원)", "created_at": "생성 일시",
        "account_type": "계좌 유형 (rec: 모의투자, prop: 투자제안)",
        "entry_strategy": "개시 전략 태그 (S1~S5 / 순수관행 / MANUAL)",
    },
    "paper_trade_histories": {
        "id": "체결 로그 ID (PK)", "trade_date": "체결 일자 (YYYY-MM-DD)",
        "trade_type": "체결 유형 (BUY / SELL / MANUAL_BUY)", "stock_code": "종목 코드",
        "stock_name": "종목명", "price": "체결 단가 (원)", "quantity": "체결 수량",
        "total_amount": "총 체결 금액", "realized_pnl": "실현 손익금 (원)", "created_at": "생성 일시",
        "account_type": "계좌 유형 (rec: 모의투자, prop: 투자제안)",
        "entry_strategy": "개시 전략 태그 (매수행)",
    },
    "strategy_daily_equity": {
        "id": "고유 ID (PK)", "combo_id": "전략 조합 ID", "trade_date": "거래 일자 (YYYYMMDD)",
        "equity_amount": "당일 총 평가자산 (원)", "created_at": "생성일시",
    },
    "strategy_leaderboard": {
        "combo_id": "전략 조합 ID (PK)", "combo_name": "전략 조합 명칭 (예: S5, S1+S2)",
        "final_capital": "최종 자산 (원)", "total_return_pct": "누적 수익률 (%)",
        "win_rate_pct": "매매 승률 (%)", "mdd_pct": "최대 낙폭 MDD (%)",
        "total_trades": "총 거래 횟수", "updated_at": "수정일시",
    },
    "strategy_trade_logs": {
        "id": "로그 ID (PK)", "combo_id": "전략 조합 ID", "trade_date": "매매 일자",
        "symbol": "종목코드", "name": "종목명", "trade_type": "매매구분 (BUY/SELL)",
        "holding_days": "보유 일수", "shares": "체결 수량", "unit_price": "체결 단가 (원)",
        "total_amount": "총 거래금액 (원)", "equity_after_trade": "거래 후 평가자산",
        "cum_return_pct": "누적 수익률 (%)", "profit_pct": "손익률 (%)",
        "profit_krw": "손익금액 (원)", "prob_up": "AI 상승확률 (%)",
        "strategy_tag": "전략 태그 (S1~S5)", "created_at": "등록일시",
        "slot_no": "포트폴리오 슬롯 번호 (1~3)",
    },
    "sync_logs": {
        "id": "로그 ID (PK)", "sync_date": "동기화 처리 일자 (YYYYMMDD)",
        "total_count": "전체 처리 건수", "kospi_count": "코스피 처리 건수",
        "kosdaq_count": "코스닥 처리 건수", "etf_count": "ETF 처리 건수",
        "etn_count": "ETN 처리 건수", "status": "실행 상태 (SUCCESS / FAILED)",
        "elapsed_seconds": "소요 시간 (초)", "elapsed_time_str": "소요 시간 표현 (시/분/초)",
        "created_at": "등록일시",
    },
    "target_stocks": {
        "symbol": "타깃 종목코드 (PK)", "created_at": "등록일시",
    },
}

HEAD_FILL = PatternFill("solid", fgColor="305496")
HEAD_FONT = Font(bold=True, color="FFFFFF", name="맑은 고딕")
BASE_FONT = Font(name="맑은 고딕")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center")


def style_header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncol)}{ws.max_row}"


def apply_body(ws, ncol, widths):
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ncol):
        for cell in row:
            cell.font = BASE_FONT
            cell.border = BORDER
            cell.alignment = WRAP
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


con = sqlite3.connect(DB)
tables = [r[0] for r in con.execute(
    "SELECT name FROM sqlite_master WHERE type='table' ORDER BY name")]

wb = Workbook()

# ---- 시트 1: 테이블 목록 ----
ws1 = wb.active
ws1.title = "테이블 목록"
ws1.append(["번호", "테이블명", "설명", "행 수", "컬럼 수", "기본키(PK)", "인덱스"])
for idx, t in enumerate(tables, start=1):
    cols = list(con.execute(f'PRAGMA table_info("{t}")'))
    pk = ", ".join([c[1] for c in sorted(cols, key=lambda x: x[5]) if c[5]])
    ncnt = len(cols)
    rcnt = con.execute(f'SELECT COUNT(*) FROM "{t}"').fetchone()[0]
    idxs = ", ".join([r[1] for r in con.execute(f'PRAGMA index_list("{t}")')
                      if not r[1].startswith("sqlite_autoindex")])
    ws1.append([idx, t, TABLE_DESC.get(t, ""), rcnt, ncnt, pk, idxs or "-"])
style_header(ws1, 7)
apply_body(ws1, 7, [6, 26, 52, 12, 9, 24, 40])

# ---- 시트 2: 컬럼 상세 ----
ws2 = wb.create_sheet("컬럼 상세")
ws2.append(["테이블명", "순서", "컬럼명", "데이터타입", "NOT NULL",
            "기본값", "PK", "설명"])
for t in tables:
    cols = list(con.execute(f'PRAGMA table_info("{t}")'))
    for c in cols:
        cid, cname, ctype, notnull, dflt, pk = c
        ws2.append([
            t, cid + 1, cname, ctype,
            "Y" if notnull else "",
            "" if dflt is None else str(dflt),
            f"PK{pk}" if pk else "",
            COL_DESC.get(t, {}).get(cname, ""),
        ])
style_header(ws2, 8)
apply_body(ws2, 8, [26, 7, 24, 16, 10, 12, 7, 56])
for c in ("B", "E", "G"):
    for cell in ws2[c][1:]:
        cell.alignment = CENTER

con.close()
OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print("saved:", OUT, "| tables:", len(tables))
